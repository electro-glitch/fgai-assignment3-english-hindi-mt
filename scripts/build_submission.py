from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

def link(p, text, url):
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r'); props=OxmlElement('w:rPr'); color=OxmlElement('w:color');color.set(qn('w:val'),'165D97');props.append(color);r.append(props)
    t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

def build(name, registration, repository):
    doc=Document(); sec=doc.sections[0]
    sec.top_margin=sec.bottom_margin=Inches(.85)
    sec.left_margin=sec.right_margin=Inches(.85)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(12)
    title=doc.styles['Title'];title.font.size=Pt(25);title.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph('Foundations of Generative AI', 'Title')
    doc.add_paragraph('ICT4403  |  Assignment 3')
    doc.add_paragraph('Submission details for the English to Hindi machine translation project comparing encoder–decoder and decoder-only language models.')
    for label,value in [('Student name',name),('Registration number',registration),('Assignment title','Machine Translation with Transformer mT5 LLaMA and Gemma'),('Autoregressive text generation system title','English to Hindi Translation with Transformer mT5 LLaMA and Gemma')]:
        p=doc.add_paragraph();p.add_run(label+'\n').bold=True;p.add_run(value)
    p=doc.add_paragraph();p.add_run('Dataset\n').bold=True;p.add_run('IIT Bombay English-Hindi Parallel Corpus\n');link(p,'Dataset source on Hugging Face','https://huggingface.co/datasets/cfilt/iitb-english-hindi')
    p=doc.add_paragraph();p.add_run('GitHub repository\n').bold=True;link(p,repository.replace('https://',''),repository)
    doc.add_paragraph('The repository contains the Python implementation and supporting configuration. Experiment results, plots, translation samples and comparative analysis accompany the completed runs.')
    for root in (doc._element,doc.styles.element):
        for el in list(root.iter(qn('w:pBdr'))): el.getparent().remove(el)
    safe_name=''.join(c for c in name if c.isalnum())
    safe_reg=''.join(c for c in registration if c.isalnum())
    path=Path(f'submission/{safe_name}_{safe_reg}_FGAI_Assignment3.docx');path.parent.mkdir(exist_ok=True);doc.save(path)
    print(path)
if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser()
    parser.add_argument('--name',required=True)
    parser.add_argument('--registration',required=True)
    parser.add_argument('--repository',required=True)
    args=parser.parse_args()
    build(args.name,args.registration,args.repository)
